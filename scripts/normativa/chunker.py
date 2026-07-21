# -*- coding: utf-8 -*-
"""
chunker.py — DOCX → chunks artículo-nivel
QUIRA Gov · Corpus Normativo Ecuatoriano · Dylus Lab © 2026

Lee un archivo DOCX y lo segmenta en chunks artículo-nivel usando patrones
regex adaptados a la estructura del derecho ecuatoriano.

Algoritmo:
  1. Extraer texto plano del DOCX (párrafos concatenados)
  2. Detectar límites de artículo con regex canónico
  3. Para artículos >450 palabras → dividir en sub-chunks con overlap
  4. El preámbulo (antes del Art. 1) → chunk especial tipo PREÁMBULO

Dependencias: python-docx
"""

from __future__ import annotations

import re
import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


# ── REGEX CANÓNICO — LÍMITES DE ARTÍCULO ─────────────────────────────────────
# Soporta:
#   Art. 226.-    Art. 226.    Art. 226 —
#   Artículo 226.-   ARTÍCULO 226.-
#   Art. 226A.-   Art. 226.1.-   (artículos innumerados o con sufijo)
#   Art. único.-  Art. innumerado.-
#
ARTICLE_RE = re.compile(
    # ANCLA (corrección 2026-07-20 · hallazgo de Javo): además del inicio de línea, se acepta el
    # marcador precedido del TÍTULO DE CAPÍTULO en el mismo párrafo — así viene maquetada la LOPC:
    #     "Capítulo Segundo\nDe la rendición de cuentas Art. 88\nDerecho ciudadano...-"
    # Exigir `^` hacía PERDER esos artículos (LOPC quedó con 77 de 103, entre ellos el 88, que es
    # el derecho ciudadano a exigir rendición de cuentas — fundamento de d09).
    # El texto previo se admite SOLO si el número cierra la línea (patrón de encabezado); una
    # referencia cruzada ("conforme al Art. 5 de esta ley") sigue el texto y por eso no matchea.
    r"(?:^|\n)(?:[^\n]{0,80}?\s)?"                # inicio de línea; opcional TÍTULO previo (caso LOPC)
    r"(?:"
    r"Art(?:í|i|Í|I)cul(?:o|a|os|as)?s?\.\s*"
    r"|ARTÍCULO\s+"
    r"|Art\.\s*"
    r")"
    r"("
    r"\d+(?:\.\d+)?[A-Za-záéíóúÁÉÍÓÚ]?"          # 226 / 226.1 / 226A  (sin tragar el título · fix Javo)
    r"|único|única|innumerado|innumerada"         # literales, NO \w+ (evita falsos matches)
    r")"
    # DISCRIMINADOR (2026-07-20 · hallazgo de Javo): un ARTÍCULO real va seguido del guión
    # dispositivo (`.-`) o CIERRA la línea (LOPC: `Art. 88\nTítulo.-`). Una REFERENCIA cruzada
    # ("del artículo 168 del Código") va seguida de espacio+palabra → el lookahead la EXCLUYE.
    # Sin esto, en leyes reformatorias (COOTAD-2026) se contaban las referencias como artículos.
    r"[ \t]*(?=[.\-–—·]|\n|$)\s*[.\-–—·]?",
    re.IGNORECASE | re.MULTILINE,
)

# DISPOSICIONES (2026-07-20 · tercer bug hallado en la auditoría): las Disposiciones
# (Transitorias/Generales/Derogatorias/Reformatorias/Finales) NO llevan "Art. N" y el chunker las
# ABSORBÍA en el último artículo previo. Ahí viven reglas operativas reales — p. ej. el umbral 65%
# de COOTAD-2026 está en la Disposición Transitoria Primera. Se reconocen como límite de chunk.
DISPOSITION_RE = re.compile(
    r"(?:^|\n)\s*"
    r"(DISPOSICI(?:Ó|O)N(?:ES)?\s+"
    r"(?:TRANSITORIA|GENERAL|DEROGATORIA|REFORMATORIA|FINAL|INTERPRETATIVA)(?:ES)?"
    r"(?:\s+(?:PRIMERA|SEGUNDA|TERCERA|CUARTA|QUINTA|SEXTA|S(?:É|E)PTIMA|OCTAVA|NOVENA|"
    r"D(?:É|E)CIMA|(?:Ú|U)NICA))?"
    r")",
    re.IGNORECASE | re.MULTILINE,
)

# Límite de palabras por chunk (alineado con CHUNK_MAX_TOKENS del RAG config)
CHUNK_MAX_WORDS = 450

# Overlap entre sub-chunks (en líneas)
CHUNK_OVERLAP_LINES = 2


@dataclass
class ArticleChunk:
    """Un chunk = un artículo (o sub-artículo si supera CHUNK_MAX_WORDS)."""
    articulo_num:   Optional[int]   # None para preámbulo
    articulo_raw:   str             # "Art. 226" tal como aparece en texto
    chunk_seq:      int             # 0=único/primero, 1,2...=continuación
    contenido:      str             # texto del chunk
    palabras:       int             # conteo de palabras
    sha256:         str             # hash SHA256 del contenido


def _sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _word_count(text: str) -> int:
    return len(text.split())


def _split_long_chunk(articulo_raw: str,
                      articulo_num: Optional[int],
                      text: str) -> list[ArticleChunk]:
    """
    Divide un artículo largo en sub-chunks de ≤CHUNK_MAX_WORDS palabras.
    Mantiene CHUNK_OVERLAP_LINES líneas de solapamiento entre chunks.
    """
    lines = text.splitlines(keepends=True)
    chunks: list[ArticleChunk] = []
    current_lines: list[str] = []
    current_words = 0
    seq = 0

    for line in lines:
        line_words = _word_count(line)
        if current_words + line_words > CHUNK_MAX_WORDS and current_lines:
            # Emitir chunk actual
            chunk_text = "".join(current_lines).strip()
            if chunk_text:
                chunks.append(ArticleChunk(
                    articulo_num=articulo_num,
                    articulo_raw=articulo_raw,
                    chunk_seq=seq,
                    contenido=chunk_text,
                    palabras=_word_count(chunk_text),
                    sha256=_sha256(chunk_text),
                ))
                seq += 1
            # Retener overlap
            overlap = current_lines[-CHUNK_OVERLAP_LINES:] if CHUNK_OVERLAP_LINES > 0 else []
            current_lines = overlap
            current_words = sum(_word_count(l) for l in overlap)

        current_lines.append(line)
        current_words += line_words

    # Emitir chunk final
    if current_lines:
        chunk_text = "".join(current_lines).strip()
        if chunk_text:
            chunks.append(ArticleChunk(
                articulo_num=articulo_num,
                articulo_raw=articulo_raw,
                chunk_seq=seq,
                contenido=chunk_text,
                palabras=_word_count(chunk_text),
                sha256=_sha256(chunk_text),
            ))

    return chunks


def _make_chunk(articulo_raw: str,
                articulo_num: Optional[int],
                text: str) -> list[ArticleChunk]:
    """Crea 1 o más chunks a partir de un artículo."""
    text = text.strip()
    if not text:
        return []
    if _word_count(text) <= CHUNK_MAX_WORDS:
        return [ArticleChunk(
            articulo_num=articulo_num,
            articulo_raw=articulo_raw,
            chunk_seq=0,
            contenido=text,
            palabras=_word_count(text),
            sha256=_sha256(text),
        )]
    return _split_long_chunk(articulo_raw, articulo_num, text)


def _parse_art_num(raw: str) -> Optional[int]:
    """Extrae el número entero del artículo a partir del match del raw."""
    try:
        # Tomar solo la parte numérica inicial
        digits = re.match(r"(\d+)", raw.strip())
        return int(digits.group(1)) if digits else None
    except (ValueError, AttributeError):
        return None


def chunk_docx(filepath: str | Path) -> list[ArticleChunk]:
    """
    Carga un DOCX y lo segmenta en chunks artículo-nivel.

    Returns:
        Lista de ArticleChunk ordenados: PREÁMBULO (si existe) + Art. 1, 2, ...
    """
    from docx import Document  # python-docx

    doc = Document(str(filepath))

    # Extraer texto completo del DOCX (párrafos, tablas)
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Incluir texto de tablas (NCI, clasificadores tienen tablas con contenido)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in parts:
                    parts.append(text)

    full_text = "\n".join(parts)

    # ── Segmentar por artículo Y disposición ──────────────────────────────────
    # Se combinan ambos tipos de límite y se ordenan por posición: así una Disposición deja de
    # absorberse en el artículo anterior (tercer bug · auditoría 2026-07-20).
    limites: list[tuple[int, str, Optional[int]]] = []
    for m in ARTICLE_RE.finditer(full_text):
        raw = m.group(1).strip()
        limites.append((m.start(), f"Art. {raw}", _parse_art_num(raw)))
    for m in DISPOSITION_RE.finditer(full_text):
        raw = re.sub(r"\s+", " ", m.group(1).strip()).title()
        limites.append((m.start(), raw, None))
    limites.sort(key=lambda x: x[0])

    if not limites:
        # Documento sin artículos ni disposiciones → todo como chunk único
        return _make_chunk("DOCUMENTO", None, full_text)

    chunks: list[ArticleChunk] = []

    # Preámbulo (texto antes del primer límite)
    preamble_text = full_text[:limites[0][0]].strip()
    if preamble_text and _word_count(preamble_text) >= 20:
        chunks.extend(_make_chunk("PREÁMBULO", None, preamble_text))

    # Artículos y disposiciones
    for i, (start, raw, num) in enumerate(limites):
        end = limites[i + 1][0] if i + 1 < len(limites) else len(full_text)
        seg_text = full_text[start:end].strip()
        chunks.extend(_make_chunk(raw, num, seg_text))

    return chunks


def chunk_docx_with_meta(filepath: str | Path,
                         meta: dict) -> list[dict]:
    """
    Wrapper que retorna dicts listos para INSERT en normativa_corpus.

    meta: entrada del MANIFEST (sigla, nombre, jerarquia, milestone, tipo, dominios)
    """
    import json
    chunks = chunk_docx(filepath)
    rows   = []
    for c in chunks:
        rows.append({
            "norma_sigla":    meta["sigla"],
            "norma_nombre":   meta["nombre"],
            "jerarquia":      meta["jerarquia"],
            "milestone_qlep": meta["milestone"],
            "tipo_documento": meta["tipo"],
            "articulo_num":   c.articulo_num,
            "articulo_raw":   c.articulo_raw,
            "chunk_seq":      c.chunk_seq,
            "contenido":      c.contenido,
            "palabras":       c.palabras,
            "dominios_quira": json.dumps(meta["dominios"], ensure_ascii=False),
            "sha256":         c.sha256,
            "archivo_nombre": Path(filepath).name,
        })
    return rows


# ── TEST RÁPIDO ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import sys, os
    if len(sys.argv) < 2:
        print("Uso: python chunker.py <ruta_archivo.docx>")
        sys.exit(1)

    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"ERROR: archivo no encontrado: {path}")
        sys.exit(1)

    chunks = chunk_docx(path)
    print(f"\nArchivo: {os.path.basename(path)}")
    print(f"Chunks detectados: {len(chunks)}\n")
    for i, c in enumerate(chunks[:20]):   # primeros 20
        print(f"  [{i+1:3d}] {c.articulo_raw:15s} seq={c.chunk_seq}  "
              f"palabras={c.palabras:4d}  "
              f"sha256={c.sha256[:12]}...")
    if len(chunks) > 20:
        print(f"  ... ({len(chunks) - 20} chunks más)")
    print()

    total_palabras = sum(c.palabras for c in chunks)
    print(f"Total palabras: {total_palabras:,}")
    print(f"Promedio por chunk: {total_palabras // max(len(chunks), 1)}")
